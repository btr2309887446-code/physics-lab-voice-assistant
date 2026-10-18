"""
对话记录导出脚本

从 SQLite 数据库导出对话记录到 Excel 文件。
所有对话记录保存在同一个 Excel 文件中。

用法：
  python export_dialogues.py                    # 导出全部对话
  python export_dialogues.py --date 2026-04-18  # 导出指定日期
  python export_dialogues.py --last 7           # 导出最近 7 天
  python export_dialogues.py --session abc123   # 导出指定会话
  python export_dialogues.py --stats            # 查看统计信息

安装依赖：
  pip install openpyxl
"""

import argparse
import os
import sqlite3
import sys
from datetime import datetime, timedelta

DB_PATH = os.path.join(os.path.dirname(__file__), "data", "conversations", "conversations.db")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "data", "conversations", "exports")


def get_connection():
    if not os.path.exists(DB_PATH):
        print(f"❌ 数据库不存在: {DB_PATH}")
        print("   请先启动服务端进行对话，数据库会自动创建。")
        sys.exit(1)
    return sqlite3.connect(DB_PATH)


def show_stats(conn):
    """显示数据库统计信息"""
    cur = conn.cursor()

    # 总体统计
    cur.execute("SELECT COUNT(*) FROM conversations")
    total_sessions = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM messages")
    total_messages = cur.fetchone()[0]
    cur.execute("SELECT COUNT(DISTINCT device_id) FROM conversations")
    total_devices = cur.fetchone()[0]

    print(f"\n📊 对话数据库统计")
    print(f"{'='*40}")
    print(f"  总会话数: {total_sessions}")
    print(f"  总消息数: {total_messages}")
    print(f"  设备数量: {total_devices}")

    if total_sessions > 0:
        cur.execute("SELECT MIN(started_at), MAX(started_at) FROM conversations")
        row = cur.fetchone()
        print(f"  时间范围: {row[0]} ~ {row[1]}")

    # 按日期统计
    cur.execute("""
        SELECT DATE(started_at) as date, COUNT(*) as session_count,
               (SELECT COUNT(*) FROM messages m WHERE DATE(m.created_at) = DATE(c.started_at)) as msg_count
        FROM conversations c
        GROUP BY DATE(started_at)
        ORDER BY date DESC
        LIMIT 14
    """)
    rows = cur.fetchall()
    if rows:
        print(f"\n📅 最近对话记录（按日期）:")
        print(f"  {'日期':<12} {'会话数':>6} {'消息数':>6}")
        print(f"  {'-'*26}")
        for row in rows:
            print(f"  {row[0]:<12} {row[1]:>6} {row[2]:>6}")

    # 按设备统计
    cur.execute("""
        SELECT device_id, COUNT(*) as session_count
        FROM conversations
        GROUP BY device_id
        ORDER BY session_count DESC
    """)
    rows = cur.fetchall()
    if rows:
        print(f"\n📱 设备会话统计:")
        print(f"  {'设备 ID':<20} {'会话数':>6}")
        print(f"  {'-'*28}")
        for row in rows:
            print(f"  {row[0]:<20} {row[1]:>6}")

    print()


def export_to_excel(conn, where_clause, params, output_file):
    """导出对话记录到 Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        print("❌ 需要 openpyxl 库。请运行: pip install openpyxl")
        sys.exit(1)

    # 查询会话
    session_query = f"""
        SELECT c.session_id, c.device_id, c.started_at, c.message_count
        FROM conversations c
        WHERE {where_clause}
        ORDER BY c.started_at
    """
    sessions = conn.cursor().execute(session_query, params).fetchall()

    if not sessions:
        print("📭 没有找到符合条件的对话记录。")
        return

    print(f"\n🔄 正在导出 {len(sessions)} 个会话...")

    wb = openpyxl.Workbook()

    # --- Sheet 1: 汇总 ---
    ws_summary = wb.active
    ws_summary.title = "会话汇总"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    summary_headers = ["#", "Session ID", "设备 ID", "开始时间", "消息数"]
    for col, h in enumerate(summary_headers, 1):
        cell = ws_summary.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for i, sess in enumerate(sessions, 1):
        for col, val in enumerate([i, sess[0], sess[1], sess[2], sess[3]], 1):
            cell = ws_summary.cell(row=i + 1, column=col, value=val)
            cell.border = thin_border

    ws_summary.column_dimensions["A"].width = 5
    ws_summary.column_dimensions["B"].width = 40
    ws_summary.column_dimensions["C"].width = 20
    ws_summary.column_dimensions["D"].width = 22
    ws_summary.column_dimensions["E"].width = 8

    # --- Sheet 2: 全部对话 ---
    ws_all = wb.create_sheet("全部对话")

    user_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    assistant_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    all_headers = ["#", "会话序号", "Session ID", "轮次", "角色", "说话人", "内容", "时间"]
    for col, h in enumerate(all_headers, 1):
        cell = ws_all.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    row_idx = 2
    global_idx = 1
    for sess_num, sess in enumerate(sessions, 1):
        session_id = sess[0]
        msgs = conn.cursor().execute("""
            SELECT turn_id, role, content, speaker, created_at
            FROM messages
            WHERE session_id = ?
            ORDER BY turn_id
        """, (session_id,)).fetchall()

        for msg in msgs:
            role_display = "👤 用户" if msg[1] == "user" else "🤖 小智"
            fill = user_fill if msg[1] == "user" else assistant_fill
            values = [global_idx, sess_num, session_id, msg[0] + 1, role_display, msg[3] or "", msg[2] or "", msg[4]]

            for col, val in enumerate(values, 1):
                cell = ws_all.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                cell.border = thin_border
                if col == 7:  # 内容列
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

            row_idx += 1
            global_idx += 1

    ws_all.column_dimensions["A"].width = 5
    ws_all.column_dimensions["B"].width = 8
    ws_all.column_dimensions["C"].width = 40
    ws_all.column_dimensions["D"].width = 6
    ws_all.column_dimensions["E"].width = 10
    ws_all.column_dimensions["F"].width = 12
    ws_all.column_dimensions["G"].width = 70
    ws_all.column_dimensions["H"].width = 20

    # 保存
    os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
    wb.save(output_file)
    print(f"✅ 导出完成: {output_file}")
    print(f"   {len(sessions)} 个会话, {global_idx - 1} 条消息")


def main():
    parser = argparse.ArgumentParser(description="导出小智对话记录到 Excel")
    parser.add_argument("--stats", action="store_true", help="查看统计信息")
    parser.add_argument("--date", help="导出指定日期 (YYYY-MM-DD)")
    parser.add_argument("--last", type=int, metavar="N", help="导出最近 N 天")
    parser.add_argument("--session", help="导出指定 session_id")
    parser.add_argument("--output", "-o", help="输出文件路径 (默认自动生成)")
    args = parser.parse_args()

    conn = get_connection()

    if args.stats:
        show_stats(conn)
        conn.close()
        return

    # 构建 WHERE 条件
    if args.session:
        where = "c.session_id = ?"
        params = (args.session,)
        label = f"session_{args.session[:8]}"
    elif args.date:
        where = "DATE(c.started_at) = ?"
        params = (args.date,)
        label = args.date
    elif args.last:
        cutoff = (datetime.now() - timedelta(days=args.last)).strftime("%Y-%m-%d")
        where = "DATE(c.started_at) >= ?"
        params = (cutoff,)
        label = f"最近{args.last}天"
    else:
        where = "1=1"
        params = ()
        label = "全部"

    # 输出文件名
    if args.output:
        output_file = args.output
    else:
        now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_file = os.path.join(OUTPUT_DIR, f"对话记录_{label}_{now_str}.xlsx")

    export_to_excel(conn, where, params, output_file)
    conn.close()


if __name__ == "__main__":
    main()
